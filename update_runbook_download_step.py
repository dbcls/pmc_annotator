from docx import Document
from docx.shared import Inches, Pt

PATH = "/Users/vsubramoniam/Code/Identifiers-Hackathon/pmc_annotator/PMC_Annotator_Sample_Runbook.docx"

def insert_before(anchor, text="", style=None, code=False):
    p = anchor.insert_paragraph_before(text, style=style)
    if code:
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(6)
        for run in p.runs:
            run.font.name = "Menlo"
            run.font.size = Pt(8.7)
    return p

doc = Document(PATH)

# Shift the existing numbered steps forward by one.
renumber = {
    "Step 0 Activate the project environment": "Step 1 Activate the project environment",
    "Step 1 Create a shard plan": "Step 2 Create a shard plan",
    "Step 2 Extract accession candidates": "Step 3 Extract accession candidates",
    "Step 3 Build candidates and mentions tables": "Step 4 Build candidates and mentions tables",
    "Step 4 Prepare T2 verification resources": "Step 5 Prepare T2 verification resources",
    "Step 5 Discover queryable TogoID graphs": "Step 6 Discover queryable TogoID graphs",
    "Step 6 Run T2 existence verification": "Step 7 Run T2 existence verification",
    "Step 7 Run T3 authoritative NCBI verification": "Step 8 Run T3 authoritative NCBI verification",
    "Step 8 Build confirmed only metrics": "Step 9 Build confirmed only metrics",
}
for p in doc.paragraphs:
    if p.text in renumber:
        p.text = renumber[p.text]
        p.style = "Heading 1"

anchor = next(p for p in doc.paragraphs if p.text == "Step 1 Activate the project environment")

insert_before(anchor, "Step 0 Download ten PMC XML articles", style="Heading 1")
insert_before(anchor, "The pipeline needs full text PMC JATS XML files, not web pages or PDFs. The commands below search PMC for ten open access or author manuscript articles likely to contain dataset accessions, then download one XML version for each article.")
insert_before(anchor, "mkdir -p /Users/vsubramoniam/Code/Identifiers-Hackathon/PMC_Data\ncd /Users/vsubramoniam/Code/Identifiers-Hackathon/PMC_Data", code=True)
insert_before(anchor, "curl -sG 'https://eutils.ncbi.nlm.nih.gov/entrez/eutils/esearch.fcgi' \\\n+  --data-urlencode 'db=pmc' \\\n+  --data-urlencode 'term=(GEO OR SRA OR transcriptome) AND (open_access[Filter] OR author_manuscript[Filter])' \\\n+  --data-urlencode 'retmax=10' \\\n+  --data-urlencode 'retmode=json' \\\n+  | jq -r '.esearchresult.idlist[]' > pmc_ids.txt\n\nwc -l pmc_ids.txt", code=True)
insert_before(anchor, "The expected result is ten numeric PMC IDs in pmc_ids.txt. The download commands use the public PMC AWS data bucket. If the aws command is unavailable on macOS, install it once with: brew install awscli.")
insert_before(anchor, "while read -r id; do\n  version_prefix=$(aws s3api list-objects-v2 \\\n+    --bucket pmc-oa-opendata \\\n+    --prefix \"PMC${id}.\" \\\n+    --delimiter \"/\" \\\n+    --query 'CommonPrefixes[0].Prefix' \\\n+    --output text \\\n+    --no-sign-request)\n\n  version=${version_prefix%/}\n  target_bucket=$(printf 'PMC%03dxxxxxx' \"$((id / 1000000))\")\n  mkdir -p \"$target_bucket\"\n\n  aws s3 cp \\\n+    \"s3://pmc-oa-opendata/${version_prefix}${version}.xml\" \\\n+    \"$target_bucket/PMC${id}.xml\" \\\n+    --no-sign-request\ndone < pmc_ids.txt\n\nfind . -name '*.xml' | wc -l", code=True)
insert_before(anchor, "Expected result: ten XML files, stored in PMC style bucket directories. This layout also supports the optional role classification stage later. PMC article licenses vary; retain the article metadata and follow the license terms if you redistribute article content.")

doc.save(PATH)
print(PATH)
